#!/usr/bin/env python3
"""
Turn one ministry's quarterly procurement report (.xlsx, as published on
gov.il) into the compact JSON the budget page reads.

WHY: BudgetKey ingests these same files but does not map the payment
column (₪7.8bn of payments reading as 0.00 in one education report,
verified) — so the paid figure comes from the ministry's own file.

THE COLUMNS (headers are irregular — match loosely, never by exact string):
  ערך ההזמנה כולל מע"מ                        → volume
  ב. חשבוניות מצטבר + מע"מ והצמדות במט"מ      → paid, CUMULATIVE to the report
  ביצוע חשבוניות  לתקופת הדוח' במטבע מקומי    → paid during the report period
                                                (only ~45% filled — not used)
  הזמנת רכש                                   → order_id
  תקנה תקציבית                                → budget code, 8 digits, no '00'

usage: parse_report.py <file.xlsx> <out-dir> [--source-url URL]
"""
import json, os, re, sys, unicodedata
import openpyxl

def norm(s):
    s = unicodedata.normalize("NFKC", str(s or ""))
    return re.sub(r"[\s‏‎]+", " ", s).strip()

def num(v):
    if v in (None, ""): return 0.0
    try: return float(str(v).replace(",", "").replace("‏", ""))
    except ValueError: return 0.0

# loose matchers, in priority order — the headers carry typos ("במט\"מ" for
# "במט\"ח"), double spaces and stray apostrophes, so exact matching is a trap
WANT = {
    "order":  [lambda h: "הזמנת רכש" in h],
    "code":   [lambda h: "תקנה תקציבית" in h],
    "vol":    [lambda h: "ערך ההזמנה" in h],
    "paid":   [lambda h: "חשבוניות" in h and "מצטבר" in h],
    "period": [lambda h: "חשבוניות" in h and "לתקופת" in h],
    "supp":   [lambda h: h.strip() == "שם הספק"],
}

def find_cols(hdr):
    out = {}
    for key, tests in WANT.items():
        for i, h in enumerate(hdr):
            hh = norm(h)
            if any(t(hh) for t in tests):
                out[key] = i
                break
    missing = [k for k in ("order", "code", "vol", "paid") if k not in out]
    if missing:
        raise SystemExit("columns not found: %s\nheaders were: %s"
                         % (missing, [norm(h) for h in hdr]))
    return out

def sheets_of(src):
    """[(sheet_name, rows_iterator), …] — decided by the BYTES, not the
       extension. Ministries published in Excel's old OLE2 format (.xls) into
       ~2020 (משרד החוץ, תיאום הפעולות בשטחים among others); openpyxl cannot
       open those, and the fetcher used to reject them as block pages. xlrd is
       imported only when an .xls actually appears, so nothing breaks where it
       is not installed and no .xls exists."""
    with open(src, "rb") as fh:
        head = fh.read(4)
    if head[:2] == b"PK":
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
        return [(name, wb[name].iter_rows(values_only=True))
                for name in wb.sheetnames]
    if head == b"\xd0\xcf\x11\xe0":
        import xlrd                     # pip install xlrd — the workflow does
        wb = xlrd.open_workbook(src)
        return [(sh.name, (tuple(sh.row_values(i)) for i in range(sh.nrows)))
                for sh in wb.sheets()]
    raise SystemExit("%s is neither xlsx nor xls (starts %s)"
                     % (os.path.basename(src), head.hex()))


VINTAGE = re.compile(r"[_/]([a-z\-]+)_([1-4])_(20\d\d)")

def report_vintage(text):
    """education_1_2025 → ("education", "2025Q1"). The quarter and year are in
       every gov.il filename; nothing else in the file records them."""
    m = VINTAGE.search(str(text).lower())
    return (m.group(1), f"{m.group(3)}Q{m.group(2)}") if m else None

def main(src, outdir, source_url=""):
    """Writes ONE FILE PER BUDGET SECTION (the first 4 digits), merging into
       whatever is already there. A ministry's report is not confined to its
       own section — the education file carries 7 rows under 0054 — and the
       page loads by section, so the shape on disk has to follow the tree,
       not the ministry."""
    by_section, full_by_section, rows, paid_rows, dupes = {}, {}, 0, 0, 0
    for sh, it in sheets_of(src):
        hdr_raw = next(it, None)
        if hdr_raw is None:
            continue                                  # an empty sheet
        cols = find_cols(hdr_raw)
        paid_hdr = norm(hdr_raw[cols["paid"]])   # the ministry's own spelling
        for r in it:
            oid = norm(r[cols["order"]])
            if not oid: continue
            code = re.sub(r"\D", "", norm(r[cols["code"]]))
            if not code: continue
            code = "00" + code if len(code) == 8 else code
            key = oid + ":" + code
            paid, vol = num(r[cols["paid"]]), num(r[cols["vol"]])
            rows += 1

            # THE WHOLE ROW, for EVERY row — including ones with no payment
            # figure (a blank paid column is not a contract that does not
            # exist; another source may hold the payment). Keys stay as the
            # ministry spelled them; an empty cell stays ABSENT, never 0.
            full = full_by_section.setdefault(code[:4], {})
            rec = {}
            for h, v in zip(hdr_raw, r):
                hk = norm(h)
                if not hk or v in (None, ""): continue
                if isinstance(v, (int, float)):
                    rec[hk] = round(float(v), 2)
                else:
                    txt = norm(v)
                    if txt and txt != "0.00" and txt != "0":
                        rec[hk] = txt
            if key not in full or paid >= num(full[key].get(paid_hdr)):
                full[key] = rec

            if paid <= 0: continue
            paid_rows += 1
            sec = by_section.setdefault(code[:4], {})
            if key in sec:
                # one order can be split across rows; the report repeats the
                # cumulative figure, so take the largest rather than summing
                dupes += 1
                if paid <= sec[key][0]: continue
            sec[key] = [round(paid, 2), round(vol, 2)]

    os.makedirs(outdir, exist_ok=True)
    for sec, orders in sorted(by_section.items()):
        path = os.path.join(outdir, sec + ".json")
        doc = {"sources": [], "orders": {}}
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                doc = json.load(f)
        if source_url and source_url not in doc["sources"]:
            doc["sources"].append(source_url)
        # record WHICH report this came from. When BudgetKey and a file
        # disagree, the newer report is the one to believe — without the
        # vintage there is no way to tell which that is.
        v = report_vintage(source_url or src)
        if v:
            doc.setdefault("reports", {})[v[0]] = v[1]
        doc["orders"].update(orders)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(doc, f, ensure_ascii=False, separators=(",", ":"))
        print(f"  {sec}.json  +{len(orders):>5}  total {len(doc['orders']):>5}")
    for sec, recs in sorted(full_by_section.items()):
        path = os.path.join(outdir, sec + ".full.json")
        doc = {"sources": [], "orders": {}}
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                doc = json.load(f)
        if source_url and source_url not in doc["sources"]:
            doc["sources"].append(source_url)
        doc["orders"].update(recs)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(doc, f, ensure_ascii=False, separators=(",", ":"))
        print(f"  {sec}.full.json  {len(doc['orders']):>5} רשומות מלאות")

    # a manifest, so the page can tell "this ministry has no file yet" (normal)
    # apart from "the data folder never got deployed" (a bug that otherwise
    # shows up as dashes everywhere and no error at all)
    secs = sorted(f[:-5] for f in os.listdir(outdir)
                  if f.endswith(".json") and f != "index.json"
                  and not f.endswith(".full.json"))
    with open(os.path.join(outdir, "index.json"), "w", encoding="utf-8") as f:
        json.dump({"sections": secs}, f, ensure_ascii=False, separators=(",", ":"))
    print(f"  index.json  sections: {', '.join(secs)}")
    print(f"{os.path.basename(src)}: rows {rows}, with a paid figure {paid_rows}, "
          f"without one {rows - paid_rows} (kept), repeated keys {dupes}")

if __name__ == "__main__":
    a = sys.argv[1:]
    url = ""
    if "--source-url" in a:
        i = a.index("--source-url"); url = a[i + 1]; a = a[:i] + a[i + 2:]
    main(a[0], a[1], url)
